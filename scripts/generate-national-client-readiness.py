#!/usr/bin/env python3
"""Build the municipality intelligence contract from the validated national workbooks.

The script is deliberately deterministic and fail-closed.  It emits four idempotent
data migrations, split by department, so the database never needs one oversized
request.  UUIDs are resolved inside PostgreSQL from the canonical municipality code.
"""

from __future__ import annotations

import argparse
import json
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


AGE_KEYS = (
    "18_25", "26_30", "31_35", "36_40", "41_45", "46_50",
    "51_55", "56_60", "61_65", "66_70", "70_plus",
)


def municipality_code(value: Any) -> str:
    if value is None:
        raise ValueError("municipality_code vacío")
    if isinstance(value, float):
        value = int(value)
    code = str(value).strip().replace(".0", "").zfill(4)
    if len(code) != 4 or not code.isdigit():
        raise ValueError(f"municipality_code inválido: {value!r}")
    return code


def integer(value: Any) -> int:
    if value in (None, "", "-"):
        return 0
    return int(round(float(value)))


def nullable_integer(value: Any) -> int | None:
    return None if value in (None, "", "-") else integer(value)


def text(value: Any) -> str | None:
    if value is None:
        return None
    value = str(value).strip()
    return value or None


def normalized_rows(path: Path, sheet: str, header_token: str) -> list[dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[sheet]
    header: list[str] | None = None
    rows: list[dict[str, Any]] = []
    for raw in worksheet.iter_rows(values_only=True):
        values = list(raw)
        if header is None:
            if any(str(value).strip() == header_token for value in values if value is not None):
                header = [str(value).strip() if value is not None else f"column_{index}" for index, value in enumerate(values)]
            continue
        if not any(value is not None for value in values):
            continue
        padded = values + [None] * (len(header) - len(values))
        rows.append(dict(zip(header, padded)))
    if header is None:
        raise ValueError(f"No se encontró {header_token!r} en {path.name}/{sheet}")
    return rows


def active_profiles(path: Path) -> dict[str, dict[str, Any]]:
    rows = normalized_rows(path, "TSE", "CODDEP")
    profiles: dict[str, dict[str, Any]] = {}
    for row in rows:
        source_department = integer(row["CODDEP"])
        if source_department == 25:  # Residentes en el extranjero no es un municipio.
            continue
        # El catálogo TSE enumera Sacatepéquez, Chimaltenango y El Progreso
        # como 02, 03 y 04; el código territorial oficial los identifica 03,
        # 04 y 02. Desde Escuintla en adelante ambos órdenes coinciden.
        canonical_department = {2: 3, 3: 4, 4: 2}.get(source_department, source_department)
        code = f"{canonical_department:02d}{integer(row['CODMUN']):02d}"
        total = integer(row["TOTAL"])
        women = integer(row["TOTAL_MUJERES"])
        men = integer(row["TOTAL_HOMBRES"])
        age_total = dict(zip(AGE_KEYS, (integer(row[key]) for key in (
            "EDAD_18a25", "EDAD_26a30", "EDAD_31a35", "EDAD_36a40", "EDAD_41a45",
            "EDAD_46a50", "EDAD_51a55", "EDAD_56a60", "EDAD_61a65", "EDAD_66a70",
            "EDAD_MAYOROIGUAL70",
        ))))
        age_women = dict(zip(AGE_KEYS, (integer(row[key]) for key in (
            "MUJERES_EDAD_18a25", "MUJERES_EDAD_26a30", "MUJERES_EDAD_31a35",
            "MUJERES_EDAD_36a40", "MUJERES_EDAD_41a45", "MUJERES_EDAD_46a50",
            "MUJERES_EDAD_51a55", "MUJERES_EDAD_56a60", "MUJERES_EDAD_61a65",
            "MUJERES_EDAD_66A70", "MUJERES_EDAD_MAYOROIGUAL70",
        ))))
        age_men = dict(zip(AGE_KEYS, (integer(row[key]) for key in (
            "HOMBRES_EDAD_18a25", "HOMBRES_EDAD_26a30", "HOMBRES_EDAD_31a35",
            "HOMBRES_EDAD_36a40", "HOMBRES_EDAD_41a45", "HOMBRES_EDAD_46a50",
            "HOMBRES_EDAD_51a55", "HOMBRES_EDAD_56a60", "HOMBRES_EDAD_61a65",
            "HOMBRES_EDAD_66A70", "HOMBRES_EDAD_MAYOROIGUAL70",
        ))))
        if women + men != total:
            raise ValueError(f"{code}: sexo no reconcilia con total")
        if sum(age_total.values()) != total:
            raise ValueError(f"{code}: edades no reconcilian con total")
        if any(age_women[key] + age_men[key] != age_total[key] for key in AGE_KEYS):
            raise ValueError(f"{code}: edades por sexo no reconcilian")
        women_literate = integer(row["MUJERES_ALFABETAS"])
        women_illiterate = integer(row["MUJERES_ANALFABETAS"])
        men_literate = integer(row["HOMBRES_ALFABETAS"])
        men_illiterate = integer(row["HOMBRES_ANALFABETAS"])
        if women_literate + women_illiterate != women or men_literate + men_illiterate != men:
            raise ValueError(f"{code}: alfabetismo no reconcilia por sexo")
        profiles[code] = {
            "total_active": total,
            "women_active": women,
            "men_active": men,
            "women_literate": women_literate,
            "women_illiterate": women_illiterate,
            "men_literate": men_literate,
            "men_illiterate": men_illiterate,
            "age_total": age_total,
            "age_women": age_women,
            "age_men": age_men,
        }
    return profiles


def census_profiles(path: Path) -> dict[str, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook["A1_2"]
    profiles: dict[str, dict[str, Any]] = {}
    for index, raw in enumerate(worksheet.iter_rows(values_only=True), 1):
        if index < 7 or raw[2] in (None, ""):
            continue
        code = municipality_code(raw[2])
        total = integer(raw[4])
        men = integer(raw[5])
        women = integer(raw[6])
        urban = integer(raw[28])
        rural = integer(raw[29])
        if men + women != total or urban + rural != total:
            raise ValueError(f"{code}: Censo 2018 no reconcilia")
        profiles[code] = {
            "population_total": total,
            "population_men": men,
            "population_women": women,
            "urban": urban,
            "rural": rural,
            "urban_share": urban / total if total else 0,
            "rural_share": rural / total if total else 0,
            "age_groups": {
                label: integer(raw[column])
                for column, label in enumerate((
                    "0_4", "5_9", "10_14", "15_19", "20_24", "25_29", "30_34",
                    "35_39", "40_44", "45_49", "50_54", "55_59", "60_64", "65_69",
                    "70_74", "75_79", "80_84", "85_89", "90_94", "95_99", "100_plus",
                ), start=7)
            },
        }
    return profiles


def grouped_history(comparative: Path, historical: Path, history_2011: Path, portal: Path) -> dict[str, dict[str, Any]]:
    comparisons = normalized_rows(comparative, "COMPARATIVO_LONG", "municipality_code")
    results_rows = normalized_rows(comparative, "RESULTADOS_LONG", "year")
    trajectories_rows = normalized_rows(historical, "TRAYECTORIAS_2011_2023", "municipality_code")
    authorities_rows = normalized_rows(historical, "AUTORIDADES_2015_2019", "election_year")
    dashboard_2011 = normalized_rows(history_2011, "RESUMEN_340", "municipality_code")
    councils_2011 = normalized_rows(history_2011, "CONCEJOS", "municipality_code")
    portal_rows = normalized_rows(portal, "MUNICIPIOS_340", "municipality_code")

    comparison_by_key = {(municipality_code(row["municipality_code"]), integer(row["year"])): row for row in comparisons}
    results_by_key: dict[tuple[str, int], list[dict[str, Any]]] = defaultdict(list)
    for row in results_rows:
        results_by_key[(municipality_code(row["municipality_code"]), integer(row["year"]))].append(row)
    summary_2011 = {municipality_code(row["municipality_code"]): row for row in dashboard_2011}
    portal_by_code = {municipality_code(row["municipality_code"]): row for row in portal_rows}
    authority_by_key: dict[tuple[str, int], list[dict[str, Any]]] = defaultdict(list)
    for row in authorities_rows:
        authority_by_key[(municipality_code(row["municipality_code"]), integer(row["election_year"]))].append(row)
    council_groups_2011: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in councils_2011:
        council_groups_2011[municipality_code(row["municipality_code"])].append({
            "party": text(row["organization"]), "seats": integer(row["council_seats"]),
        })
    trajectories_by_code: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in trajectories_rows:
        trajectories_by_code[municipality_code(row["municipality_code"])].append({
            "name": text(row["person"]),
            "years": [integer(item) for item in str(row["years"]).split(",") if item.strip()],
            "elections": integer(row["elections"]),
            "route": text(row["trajectory"]),
            "match_method": text(row.get("match_method")),
            "caution": text(row["caution"]),
        })

    output: dict[str, dict[str, Any]] = {}
    for code in sorted(portal_by_code):
        elections: list[dict[str, Any]] = []
        for year in (2011, 2015, 2019, 2023):
            comparison = comparison_by_key[(code, year)]
            ranked_raw = sorted(results_by_key[(code, year)], key=lambda row: integer(row["votes"]), reverse=True)
            party_votes = sum(integer(row["votes"]) for row in ranked_raw)
            ranked = [{
                "rank": rank,
                "party": text(row["organization"]),
                "candidate": text(row["candidate"]),
                "votes": integer(row["votes"]),
                "share": integer(row["votes"]) / party_votes if party_votes else 0,
                "source_id": text(row["source_id"]),
            } for rank, row in enumerate(ranked_raw, 1)]
            summary = summary_2011.get(code, {}) if year == 2011 else {}
            portal_row = portal_by_code[code] if year == 2023 else {}
            elected_mayor = next((
                row for row in authority_by_key[(code, year)]
                if (text(row.get("office")) or "").upper().startswith("ALCALDE")
                and (
                    not text(comparison.get("winner_party"))
                    or text(row.get("organization")) == text(comparison.get("winner_party"))
                )
            ), None)
            winner_candidate = (
                text(comparison.get("winner_candidate"))
                or text(elected_mayor.get("person_name") if elected_mayor else None)
                or (ranked[0]["candidate"] if ranked else None)
            )
            elections.append({
                "year": year,
                "winner_party": text(comparison["winner_party"]),
                "winner_candidate": winner_candidate,
                "winner_votes": integer(comparison["winner_votes"]),
                "runner_up_party": text(comparison["runner_up_party"]),
                "runner_up_votes": integer(comparison["runner_up_votes"]),
                "organizations": integer(comparison["organizations"]),
                "registered_voters": nullable_integer(summary.get("registered_voters")) if year == 2011 else nullable_integer(portal_row.get("registered_voters_2023")) if year == 2023 else None,
                "votes_cast": nullable_integer(summary.get("votes_cast")),
                "null_votes": nullable_integer(summary.get("null_votes")),
                "blank_votes": nullable_integer(summary.get("blank_votes")),
                "party_votes": party_votes,
                "margin_votes": integer(comparison["winner_votes"]) - integer(comparison["runner_up_votes"]),
                "results": ranked,
                "status": text(comparison["status"]),
            })

        councils: list[dict[str, Any]] = [{
            "year": 2011,
            "total": sum(group["seats"] for group in council_groups_2011[code]),
            "groups": council_groups_2011[code],
            "members": [],
            "detail_status": "SEAT_DISTRIBUTION_ONLY",
        }]
        for year in (2015, 2019):
            members = [{
                "office": text(row["office"]), "name": text(row["person_name"]),
                "party": text(row["organization"]), "source_page": nullable_integer(row["pdf_page"]),
            } for row in authority_by_key[(code, year)]]
            groups = Counter(member["party"] for member in members if member["party"])
            councils.append({
                "year": year,
                "total": len(members),
                "groups": [{"party": party, "seats": seats} for party, seats in groups.most_common()],
                "members": members,
                "detail_status": "FULL_MEMBERS",
            })
        output[code] = {
            "elections": elections,
            "councils": councils,
            "trajectories": trajectories_by_code[code],
        }
    return output


def community_catalog(path: Path) -> dict[str, dict[str, Any]]:
    rows = normalized_rows(path, "COMUNIDADES_LONG", "municipality_code")
    by_code: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        code = municipality_code(row["municipality_code"])
        by_code[code].append({
            "name": text(row["community_name"]),
            "category": text(row["category"]),
            "group": text(row["group_name"]),
            "group_code": text(row["group_code"]),
            "scope": text(row["scope"]),
            "zone": text(row["zone"]),
            "reference": text(row["reference"]),
            "marked_cem_center": bool(row["marked_cem_center"]),
            "source_page": nullable_integer(row["pdf_page"]),
        })
    output: dict[str, dict[str, Any]] = {}
    for code, records in by_code.items():
        group_counts = Counter((row["group_code"], row["group"], row["scope"]) for row in records)
        categories = Counter(row["category"] or "SIN_CATEGORIA" for row in records)
        output[code] = {
            "summary": {
                "records": len(records),
                "urban": sum(1 for row in records if row["scope"] == "URBANO"),
                "rural": sum(1 for row in records if row["scope"] == "RURAL"),
                "groups": [{"code": key[0], "name": key[1], "scope": key[2], "records": count} for key, count in sorted(group_counts.items())],
                "categories": [{"name": name, "count": count} for name, count in categories.most_common()],
            },
            "records": records,
        }
    return output


def center_catalog(path: Path) -> dict[str, list[dict[str, Any]]]:
    rows = normalized_rows(path, "CENTROS_MAPA", "location_id")
    by_code: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in rows:
        source_code = municipality_code(row["municipality_code"])
        department_name = (text(row["department"]) or "").upper()
        canonical_prefix = {
            "EL PROGRESO": "02",
            "SACATEPÉQUEZ": "03",
            "SACATEPEQUEZ": "03",
            "CHIMALTENANGO": "04",
        }.get(department_name, source_code[:2])
        code = canonical_prefix + source_code[2:]
        by_code[code].append({
            "location_id": text(row["location_id"]),
            "center_sequences": text(row["center_sequences"]),
            "communities": text(row["communities"]),
            "name": text(row["center_name"]),
            "address": text(row["address"]),
            "zone": nullable_integer(row["zone"]),
            "institution": text(row["institution"]),
            "registered_voters": nullable_integer(row["registered_voters"]),
            "jrv": nullable_integer(row["jrv"]),
            "jrv_ranges": text(row["jrv_ranges"]),
            "latitude": row["latitude"],
            "longitude": row["longitude"],
            "geolocation_status": text(row["geolocation_status"]),
        })
    return by_code


def sql_literal(value: Any) -> str:
    return "'" + str(value).replace("'", "''") + "'"


def json_literal(value: Any) -> str:
    return sql_literal(json.dumps(value, ensure_ascii=False, separators=(",", ":"))) + "::jsonb"


def assert_complete(name: str, mapping: dict[str, Any], codes: set[str]) -> None:
    missing = sorted(codes - set(mapping))
    extra = sorted(set(mapping) - codes)
    if missing or extra:
        raise ValueError(f"{name}: missing={missing[:8]} extra={extra[:8]}")


def render_batch(codes: Iterable[str], profiles: dict[str, dict[str, Any]], output: Path) -> None:
    codes = list(codes)
    profile_values = []
    active_values = []
    for code in codes:
        profile = profiles[code]
        profile_values.append(f"({sql_literal(code)},{json_literal(profile)},{sql_literal('INTELLIGENCE_READY')},ARRAY[]::text[])")
        active = profile["active_voter_profile"]
        active_values.append(
            "(" + ",".join((
                sql_literal(code), str(active["total_active"]), str(active["women_active"]), str(active["men_active"]),
                str(active["women_literate"]), str(active["women_illiterate"]), str(active["men_literate"]), str(active["men_illiterate"]),
                json_literal(active["age_total"]), json_literal(active["age_women"]), json_literal(active["age_men"]),
            )) + ")"
        )
    sql = f"""-- Generated by scripts/generate-national-client-readiness.py. Do not hand-edit.
with input(municipality_code, profile, readiness_status, missing_requirements) as (
  values\n    {',\n    '.join(profile_values)}
)
insert into data_vault.municipality_intelligence_profiles_v1(
  country_code, municipality_id, profile, readiness_status, missing_requirements,
  source_manifest, generated_at, updated_at
)
select 'GT', m.id, i.profile, i.readiness_status, i.missing_requirements,
       i.profile->'source_manifest', now(), now()
from input i
join public.municipalities m on m.country_code='GT' and m.municipality_code=i.municipality_code
on conflict (municipality_id) do update set
  profile=excluded.profile,
  readiness_status=excluded.readiness_status,
  missing_requirements=excluded.missing_requirements,
  source_manifest=excluded.source_manifest,
  generated_at=excluded.generated_at,
  updated_at=now();

with input(
  municipality_code,total_active,women_active,men_active,women_literate,women_illiterate,
  men_literate,men_illiterate,age_total,age_women,age_men
) as (
  values\n    {',\n    '.join(active_values)}
)
insert into data_vault.tse_active_voter_profiles_2026(
  country_code,municipality_id,cutoff_at,total_active,women_active,men_active,
  women_literate,women_illiterate,men_literate,men_illiterate,age_total,age_women,age_men,
  source_id,source_label,source_status,updated_at
)
select 'GT',m.id,'2026-07-12 22:01:03-06'::timestamptz,i.total_active,i.women_active,i.men_active,
       i.women_literate,i.women_illiterate,i.men_literate,i.men_illiterate,i.age_total,i.age_women,i.age_men,
       'GT-TSE-EMPADRONADOS-ACTIVOS-2026-001','TSE · Ciudadanos empadronados activos · corte 12 julio 2026','VALIDATED',now()
from input i
join public.municipalities m on m.country_code='GT' and m.municipality_code=i.municipality_code
on conflict (municipality_id,cutoff_at,source_id) do update set
  total_active=excluded.total_active,women_active=excluded.women_active,men_active=excluded.men_active,
  women_literate=excluded.women_literate,women_illiterate=excluded.women_illiterate,
  men_literate=excluded.men_literate,men_illiterate=excluded.men_illiterate,
  age_total=excluded.age_total,age_women=excluded.age_women,age_men=excluded.age_men,
  source_label=excluded.source_label,source_status=excluded.source_status,updated_at=now();
"""
    output.write_text(sql, encoding="utf-8")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--active", type=Path, required=True)
    parser.add_argument("--census", type=Path, required=True)
    parser.add_argument("--comparative", type=Path, required=True)
    parser.add_argument("--historical", type=Path, required=True)
    parser.add_argument("--history-2011", type=Path, required=True)
    parser.add_argument("--portal", type=Path, required=True)
    parser.add_argument("--output-dir", type=Path, required=True)
    args = parser.parse_args()

    active = active_profiles(args.active)
    census = census_profiles(args.census)
    history = grouped_history(args.comparative, args.historical, args.history_2011, args.portal)
    communities = community_catalog(args.portal)
    centers = center_catalog(args.portal)
    codes = set(active)
    if len(codes) != 340:
        raise ValueError(f"Padrón activo: {len(codes)}/340")
    for name, mapping in (("Censo", census), ("Histórico", history), ("Comunidades", communities), ("Centros", centers)):
        assert_complete(name, mapping, codes)

    source_manifest = {
        "active_voters": "GT-TSE-EMPADRONADOS-ACTIVOS-2026-001",
        "census": "INE-CENSO-2018-A1.2",
        "electoral_history": ["GT-TSE-MEM-2011-001", "GT-TSE-MEM-2015-001", "GT-TSE-MEM-2019-001", "GT-TSE-MEM-2023-001"],
        "community_catalog": "GT-TSE-CEM-2023-001",
        "center_directory": "GT-TSE-CENTROS-JRV-2023-001",
    }
    profiles = {
        code: {
            "municipality_code": code,
            "active_voter_profile": active[code],
            "census_2018": census[code],
            "electoral_history": history[code],
            "community_catalog": communities[code],
            "voting_centers": centers[code],
            "source_manifest": source_manifest,
        }
        for code in sorted(codes)
    }

    args.output_dir.mkdir(parents=True, exist_ok=True)
    departments = [(1, 6), (7, 12), (13, 17), (18, 22)]
    for batch, (first, last) in enumerate(departments, 1):
        batch_codes = [code for code in sorted(codes) if first <= int(code[:2]) <= last]
        render_batch(batch_codes, profiles, args.output_dir / f"20260919090{batch}00_load_national_intelligence_profiles_batch_{batch}.sql")

    report = {
        "status": "PASS",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "municipalities": len(codes),
        "active_profiles": len(active),
        "census_profiles": len(census),
        "electoral_histories": len(history),
        "community_records": sum(len(value["records"]) for value in communities.values()),
        "voting_centers": sum(len(value) for value in centers.values()),
        "sibinal": {
            "active": active["1208"]["total_active"],
            "age_bands": len(active["1208"]["age_total"]),
            "census_total": census["1208"]["population_total"],
            "urban": census["1208"]["urban"],
            "rural": census["1208"]["rural"],
            "history_years": [item["year"] for item in history["1208"]["elections"]],
            "communities": communities["1208"]["summary"]["records"],
            "centers": len(centers["1208"]),
        },
    }
    (args.output_dir / "national-client-readiness-report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
